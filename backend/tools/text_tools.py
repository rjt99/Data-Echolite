import json
import pandas as pd


def read_docx(file_path: str) -> list:
    from docx import Document
    doc = Document(file_path)
    paragraphs = []
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            paragraphs.append({"index": i, "style": para.style.name, "text": para.text.strip()})
    return paragraphs


def read_pptx(file_path: str) -> list:
    from pptx import Presentation
    prs = Presentation(file_path)
    slides = []
    for slide_num, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            slides.append({"slide": slide_num + 1, "text": "\n".join(texts)})
    return slides


def read_csv_or_excel(file_path: str) -> dict:
    ext = file_path.rsplit(".", 1)[-1].lower()
    df = pd.read_csv(file_path) if ext == "csv" else pd.read_excel(file_path)
    return {
        "columns": list(df.columns),
        "row_count": len(df),
        "sample": df.head(5).to_dict(orient="records"),
        "dtypes": {col: str(dtype) for col, dtype in df.dtypes.items()},
        "null_counts": df.isnull().sum().to_dict(),
        "text_repr": df.to_string(max_rows=50, max_cols=20)
    }


def read_json(file_path: str) -> dict:
    with open(file_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    return {"content": data, "text_repr": json.dumps(data, indent=2)[:5000]}


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> list:
    if len(text) <= chunk_size:
        return [text]
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        last_period = chunk.rfind(". ")
        if last_period > chunk_size * 0.5:
            end = start + last_period + 1
            chunk = text[start:end]
        chunks.append(chunk.strip())
        start = end - overlap
    return [c for c in chunks if c.strip()]
